import io
import openpyxl
import pandas as pd
from src.data_collection.data_maps import DataFrameMap
from streamlit.typing import UploadedFile
from unidecode import unidecode

def scan_mlb_probables(file: UploadedFile) -> DataFrameMap:
    file_bytes = io.BytesIO(file.read())

    workbook = openpyxl.load_workbook(
        filename=file_bytes,
        data_only=True
    )

    sheet = workbook[workbook.sheetnames[0]]

    rows = list(sheet.iter_rows(values_only=False))
    if not rows:
        return DataFrameMap(
            error=f"ERROR (Excel-Parser): Failed to scan ROWS of sheet ({sheet})",
            content=None
        )
    
    df = pd.DataFrame()

    for row in rows:
        cells = [cell for cell in row]
        new_row_idx = df.shape[0]

        for cell_idx, cell in enumerate(cells):
            df.loc[new_row_idx, cell_idx] = cell.value

    return df