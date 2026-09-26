import datetime
import io
import openpyxl
import pandas as pd
from src.data_collection.data_maps import DataFrameMap
from streamlit.typing import UploadedFile
from unidecode import unidecode

def scan_pitchers_data(
    file: UploadedFile,
    timestamp: datetime.datetime
) -> DataFrameMap:

    file_bytes = io.BytesIO(file.read())

    workbook = openpyxl.load_workbook(
        filename=file_bytes,
        data_only=True
    )

    sheet = workbook[workbook.sheetnames[0]]

    rows = list(sheet.iter_rows(values_only=False))
    if not rows:
        return DataFrameMap(
            error=f"ERROR (Excel-Parser): Failed to scan ROWS of sheet ({sheet})",
            content=None
        )

    headers = [cell.value.strip().upper() for cell in rows[0]]
    if not headers:
        return DataFrameMap(
            error=f"ERROR (Excel-Parser): Failed to scan HEADERS of sheet ({sheet})",
            content=None
        )

    pitchers_df = pd.DataFrame()

    for row_idx, row in enumerate(rows):
        if 0 == row_idx:
            continue

        cells = [cell for cell in row]

        new_row_idx = pitchers_df.shape[0]

        for header, cell in zip(headers, cells):
            pitchers_df.loc[new_row_idx, header] = cell.value

            cell_link = cell.hyperlink

            if cell_link:
                pitchers_df.loc[new_row_idx, f"{header}_LINK"] = cell_link.target

    pitchers_df.drop(
        columns=["#", "TEAM_LINK"],
        inplace=True
    )

    pitchers_df.rename(
        columns={
            "NAME": "PLAYER",
            "NAME_LINK": "PLAYER FGID"
        },
        inplace=True
    )

    pitchers_df["PLAYER FGID"] = pitchers_df["PLAYER FGID"].apply(
        lambda link: link.split("/stats")[0].split("/")[-1]
    )

    pitchers_df["PLAYER"] = pitchers_df["PLAYER"].apply(unidecode)

    return DataFrameMap(
        error=None,
        content=pitchers_df
    )