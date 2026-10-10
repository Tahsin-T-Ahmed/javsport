import io
import openpyxl
import pandas as pd
from src.data_collection.data_maps import DataFrameMap
from streamlit.delta_generator import DeltaGenerator
from streamlit.typing import UploadedFile
from unidecode import unidecode

def scan_mlb_roster(
    file: UploadedFile,
    progress_bar: DeltaGenerator | None = None,
    **kwargs
) -> DataFrameMap:
    file_bytes = io.BytesIO(file.read())

    workbook = openpyxl.load_workbook(
        filename=file_bytes,
        data_only=True
    )

    sheet = workbook[workbook.sheetnames[0]]

    rows = list(sheet.iter_rows(values_only=False))
    if not rows:
        return DataFrameMap(
            error=f"ERROR (MLB Roster-Scanner): Failed to scan ROWS of file",
            content=None
        )

    roster_df = pd.DataFrame()

    n_rows = len(rows)

    for row_idx, row in enumerate(rows):
        if progress_bar:
            progress_bar.progress(
                value=row_idx / n_rows,
                text=f"Scanning row {row_idx+1} of {n_rows} ({int(row_idx/n_rows*100)}%)..."
            )

        cells = [cell for cell in row]
        if not cells:
            return DataFrameMap(
                error=f"ERROR (MLB Roster-Scanner): Failed to scan CELLS of row #{row_idx+1} in file",
                content=None
            )

        for cell_idx, cell in enumerate(cells):
            if not cell.hyperlink:
                continue

            new_row_idx = roster_df.shape[0]

            roster_df.loc[new_row_idx, "ROLE"] = str(cells[cell_idx-2].value).strip()
            roster_df.loc[new_row_idx, "POSITION"] = str(cells[cell_idx-1].value).strip()
            roster_df.loc[new_row_idx, "PLAYER NAME"] = str(unidecode(cell.value)).strip()
            roster_df.loc[new_row_idx, "PLAYER FGID"] = str(cell.hyperlink.target.split("/stats")[0].split("/")[-1]).strip()

    if progress_bar:
        progress_bar.status(
            label=":green[Complete! :material/check:]",
            state="complete"
        )

    return DataFrameMap(
        error=None,
        content=roster_df
    )