import io
import openpyxl
import pandas as pd
from src.data_collection.data_maps import DataFrameMap
from streamlit.typing import UploadedFile
from unidecode import unidecode

import streamlit as st

def scan_mlb_starting_pitchers(
    file: UploadedFile,
    **kwargs
) -> DataFrameMap:
    file_bytes = io.BytesIO(file.read())

    workbook = openpyxl.load_workbook(
        filename=file_bytes,
        data_only=True
    )

    sheet = workbook[workbook.sheetnames[0]]

    rows = list(sheet.iter_rows(values_only=False))
    if not rows:
        return DataFrameMap(
            error=f"ERROR (MLB Roster-Scanner): Failed to scan ROWS of file",
            content=None
        )

    starting_pitchers = pd.DataFrame()

    for row_idx, row in enumerate(rows):
        cells = [cell for cell in row]
        if not cells:
            return DataFrameMap(
                error=f"ERROR (MLB Roster-Scanner): Failed to scan CELLS of row #{row_idx+1} in file",
                content=None
            )        

        for cell_idx, cell in enumerate(cells):
            if "SP" == cell.value and "SP" in cells[cell_idx-1].value:
                pitcher_cell = cells[cell_idx+1]
                pitcher_name = unidecode(pitcher_cell.value)
                pitcher_link = None

                if pitcher_cell.hyperlink:
                    pitcher_link = pitcher_cell.hyperlink.target

                pitcher_fgid = pitcher_link.split("/stats")[0].split("/")[-1]

                new_row_idx = starting_pitchers.shape[0]
                
                starting_pitchers.loc[new_row_idx, "NAME"] = pitcher_name
                starting_pitchers.loc[new_row_idx, "FGID"] = pitcher_fgid

    return DataFrameMap(
        error=None,
        content=starting_pitchers
    )